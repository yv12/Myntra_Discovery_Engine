"""
Survey loader for Discovery Engine.
Reads survey responses from Excel, converts each response to a normalized record,
concatenates responses for classification, and preserves full structured answers
in meta (especially the elasticity question verbatim).
"""

import openpyxl
import hashlib
from datetime import datetime
from typing import Iterator, Dict, Any

def load_survey(file_path: str = "data/input/survey_responses.xlsx") -> Iterator[Dict[str, Any]]:
    """
    Yields normalized records from Survey XLSX.
    Schema:
      - record_id: sha256 of "survey:row_{index}"
      - source: "survey"
      - text: human-readable concatenation of answers
      - created_at: ISO 8601 string or None
      - meta: full structured answers verbatim, including elasticity
    """
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    
    # Collect headers
    headers = [str(ws.cell(1, col).value or "").strip() for col in range(1, ws.max_column + 1)]
    
    row_idx = 0
    for r in range(2, ws.max_row + 1):
        row_vals = [ws.cell(r, col).value for col in range(1, ws.max_column + 1)]
        if not any(v is not None and str(v).strip() != "" for v in row_vals):
            continue
            
        row_idx += 1
        
        # Build verbatim structured meta
        row_dict = {}
        for h, v in zip(headers, row_vals):
            if h:
                # Handle datetime
                if isinstance(v, datetime):
                    row_dict[h] = v.isoformat()
                else:
                    row_dict[h] = str(v).strip() if v is not None else None
                    
        # Parse timestamp
        raw_ts = row_vals[0]
        if isinstance(raw_ts, datetime):
            created_at = raw_ts.isoformat()
        elif raw_ts:
            created_at = str(raw_ts)
        else:
            created_at = None

        # Build concatenated text for classifier
        text_parts = []
        q_screen = row_dict.get("Do you have something in your wishlist you saved more than two weeks ago and still haven't bought?")
        q_often = row_dict.get("How often does something you've saved actually end up being bought?")
        q_main_reason = row_dict.get("What's the main reason things stay in your wishlist?")
        q_saved_when = row_dict.get("How long ago did you save it?")
        q_why_saved = row_dict.get("Why did you save it instead of buying it then?")
        q_since = row_dict.get("What's happened with it since?")
        q_fit_conf = row_dict.get("How sure are you about this item today? [Whether the size and fit are right for me]")
        q_qual_conf = row_dict.get("How sure are you about this item today? [Whether the quality is as good as it looks]")
        q_blocker = row_dict.get("What's the ONE thing most stopping you from buying it? Pick the biggest one.")
        q_help = row_dict.get("Which ONE of these would help you the most with that?")
        q_elasticity = row_dict.get("If that were sorted out tomorrow, what would you honestly do?")
        q_decay = row_dict.get("What made you go off it?")

        if q_main_reason:
            text_parts.append(f"Main reason in wishlist: {q_main_reason}")
        if q_why_saved:
            text_parts.append(f"Why saved instead of bought: {q_why_saved}")
        if q_since:
            text_parts.append(f"Status since saving: {q_since}")
        if q_blocker:
            text_parts.append(f"Biggest purchase blocker: {q_blocker}")
        if q_help:
            text_parts.append(f"Helpful resolution: {q_help}")
        if q_elasticity:
            text_parts.append(f"Action if resolved: {q_elasticity}")
        if q_decay:
            text_parts.append(f"Why lost interest: {q_decay}")
        if q_fit_conf:
            text_parts.append(f"Fit confidence rating: {q_fit_conf}")
        if q_qual_conf:
            text_parts.append(f"Quality confidence rating: {q_qual_conf}")

        text = "\n".join(text_parts).strip()
        if not text:
            text = f"Wishlist item response {row_idx}"

        record_id = hashlib.sha256(f"survey:row_{row_idx}".encode("utf-8")).hexdigest()
        
        # Meta flags
        meta = dict(row_dict)
        meta["row_index"] = row_idx
        meta["screened_in"] = (q_screen == "Yes")

        yield {
            "record_id": record_id,
            "source": "survey",
            "text": text,
            "created_at": created_at,
            "meta": meta
        }
